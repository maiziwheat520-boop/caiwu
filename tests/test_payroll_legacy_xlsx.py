# ruff: noqa: RUF001

from __future__ import annotations

from pathlib import Path

import openpyxl

from ledgerbridge.payroll_legacy_xlsx import import_calculated_payroll
from ledgerbridge.payroll_module import PaymentChannel, validate_draft


def _write_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    release = workbook.create_sheet("工资发放表")
    release.append(["标题"])
    release.append(["姓名", "工作地点", "实领人", "卡号", "总工资（元）", "备注"])
    release.append(["员工甲", "星汇", "收款人甲", "622200001234", 5500, "绩效500 应发5500"])

    detail = workbook.create_sheet("工资明细表")
    detail.append(["标题"])
    detail_headers = [
        "姓名",
        "岗位",
        "工作地点",
        "基本工资（元）",
        "岗位补贴（元）",
        "卫生补贴（元）",
        "夜班（元）",
        "技术补贴（元）",
        "出行补贴（元）",
        "通讯补贴（元）",
        "住宿（元）",
        "餐补（元）",
        "化妆补贴（元）",
        "全勤（元）",
        "请假（元）",
        "加班（元）",
        "考核（元）",
        "考核说明",
        "工资（元）",
        "社保（元）",
        "代扣（元）",
        "好评（元）",
        "应发",
        "特殊项金额（元）",
        "特殊项说明",
        "备注",
    ]
    detail.append(detail_headers)
    values = {header: 0 for header in detail_headers}
    values.update(
        {
            "姓名": "员工甲",
            "岗位": "前台",
            "工作地点": "星汇",
            "基本工资（元）": 5000,
            "工资（元）": 5000,
            "好评（元）": 500,
            "应发": 5500,
            "备注": "绩效500 应发5500",
        }
    )
    detail.append([values[header] for header in detail_headers])

    info = workbook.create_sheet("员工信息表")
    info.append(["标题"])
    info.append(
        ["姓名", "卡号", "工作地点", "岗位", "是否陈队", "是否现金发放", "别名", "人员类型"]
    )
    info.append(["员工甲", "622200001234", "星汇", "前台", "否", "否", "收款人甲", "正式工"])

    attendance = workbook.create_sheet("考勤表")
    attendance.append(["标题"])
    attendance.append(["姓名", "考勤天数（天）"])
    attendance.append(["员工甲", 31])
    workbook.save(path)


def test_calculated_workbook_becomes_a_valid_split_payroll_draft(tmp_path: Path) -> None:
    path = tmp_path / "26.8.xlsx"
    _write_workbook(path)

    draft = import_calculated_payroll(
        path,
        company_id="hotel-group",
        rules_version="rules-2026-08-v1",
        cash_splits_minor={"员工甲": 100_000},
    )

    assert draft.period == "2026-08"
    assert len(draft.lines) == 1
    line = draft.lines[0]
    assert line.payment_channel is PaymentChannel.MYBANK
    assert line.net_amount_minor == 550_000
    assert line.cash_amount_minor == 100_000
    assert line.bank_amount_minor == 450_000
    assert line.payee_name == "收款人甲"
    assert validate_draft(draft) == ()
