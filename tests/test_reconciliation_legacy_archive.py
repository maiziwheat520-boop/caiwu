from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
import pytest

from ledgerbridge.reconciliation_legacy_archive import (
    LegacyArchiveError,
    extract_legacy_archive,
)


def _source(path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "24.01"
    first["A1"] = "收入"
    first["B1"] = 100.25
    first["C1"] = "=B1*2"
    second = workbook.create_sheet("24.2")
    second["A1"] = "结余"
    second["B1"] = 0
    workbook.save(path)


def test_archive_preserves_source_and_formula_distinction(tmp_path: Path) -> None:
    path = tmp_path / "history.xlsx"
    _source(path)

    archive = extract_legacy_archive(path)

    assert archive.source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert [sheet.period for sheet in archive.sheets] == ["2024-01", "2024-02"]
    cells = json.loads(archive.sheets[0].cells_json)
    assert cells[0]["address"] == "A1"
    assert cells[1]["type"] == "number"
    assert cells[1]["value"] == "100.25"
    assert cells[2] == {
        "address": "C1",
        "type": "formula",
        "value": "=B1*2",
        "cached_type": "blank",
        "cached_value": None,
    }
    assert archive.sheets[1].cell_count == 2


def test_archive_rejects_duplicate_period(tmp_path: Path) -> None:
    path = tmp_path / "history.xlsx"
    _source(path)
    workbook = openpyxl.load_workbook(path)
    workbook.copy_worksheet(workbook["24.01"]).title = "24.1"
    workbook.save(path)

    with pytest.raises(LegacyArchiveError, match="同一月份"):
        extract_legacy_archive(path)


def test_archive_rejects_non_workbook(tmp_path: Path) -> None:
    path = tmp_path / "history.xlsx"
    path.write_bytes(b"not a workbook")

    with pytest.raises(LegacyArchiveError, match="有效的"):
        extract_legacy_archive(path)
