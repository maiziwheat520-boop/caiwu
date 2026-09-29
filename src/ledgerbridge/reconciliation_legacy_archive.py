"""Lossless, read-only extraction of historical reconciliation sheets.

The original workbook bytes remain the evidence. Cell snapshots are a queryable
view of its saved values, not a recalculation or a new accounting assertion.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID
from zipfile import BadZipFile, ZipFile

import openpyxl  # type: ignore[import-untyped]
from sqlalchemy import text
from sqlalchemy.orm import Session

_PERIOD = re.compile(r"^(?P<year>\d{2})\.(?P<month>1[0-2]|0?[1-9])$")
_MAX_SOURCE_BYTES = 10 * 1024 * 1024
_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
_MAX_SHEETS = 120
_MAX_CELLS = 100_000


class LegacyArchiveError(ValueError):
    """The archive cannot be imported without losing source identity or content."""


@dataclass(frozen=True, slots=True)
class ArchivedSheet:
    period: str
    sheet_name: str
    row_count: int
    cell_count: int
    content_sha256: str
    cells_json: str


@dataclass(frozen=True, slots=True)
class ArchivedWorkbook:
    source_sha256: str
    source_bytes: bytes
    original_filename: str
    sheets: tuple[ArchivedSheet, ...]


def _encoded_value(value: Any) -> tuple[str, str | None]:
    if value is None:
        return "blank", None
    if isinstance(value, bool):
        return "boolean", "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return "date", value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        return "number", str(value)
    if isinstance(value, str):
        return "text", value
    raise LegacyArchiveError("工作簿包含不支持的单元格值类型")


def extract_legacy_archive(path: str | Path) -> ArchivedWorkbook:
    """Preserve original bytes, formula text, and saved cached values separately."""

    source_path = Path(path)
    if source_path.stat().st_size > _MAX_SOURCE_BYTES:
        raise LegacyArchiveError("历史对账来源文件超过大小上限")
    return extract_legacy_archive_bytes(source_path.read_bytes(), source_path.name)


def extract_legacy_archive_bytes(source_bytes: bytes, filename: str) -> ArchivedWorkbook:
    """Read a bounded upload without writing the financial source to a temporary file."""

    if (
        Path(filename).name != filename
        or "\\" in filename
        or Path(filename).suffix.lower() != ".xlsx"
    ):
        raise LegacyArchiveError("历史对账来源必须是 .xlsx 文件")
    if not source_bytes or len(source_bytes) > _MAX_SOURCE_BYTES:
        raise LegacyArchiveError("历史对账来源文件超过大小上限")
    try:
        with ZipFile(BytesIO(source_bytes)) as package:
            infos = package.infolist()
            if len(infos) > 500 or sum(item.file_size for item in infos) > _MAX_UNCOMPRESSED_BYTES:
                raise LegacyArchiveError("工作簿解压后超过大小上限")
    except BadZipFile as exc:
        raise LegacyArchiveError("历史对账来源不是有效的 .xlsx 文件") from exc

    formulas = openpyxl.load_workbook(BytesIO(source_bytes), read_only=True, data_only=False)
    cached = openpyxl.load_workbook(BytesIO(source_bytes), read_only=True, data_only=True)
    try:
        names = [name for name in formulas.sheetnames if _PERIOD.fullmatch(name.strip())]
        if not names or len(names) > _MAX_SHEETS:
            raise LegacyArchiveError("历史对账月份工作表数量无效")
        sheets: list[ArchivedSheet] = []
        seen_periods: set[str] = set()
        total_cells = 0
        for name in names:
            match = _PERIOD.fullmatch(name.strip())
            assert match is not None
            period = f"20{match['year']}-{int(match['month']):02d}"
            if period in seen_periods:
                raise LegacyArchiveError("同一月份存在多个历史对账工作表")
            seen_periods.add(period)
            formula_sheet = formulas[name]
            cache_sheet = cached[name]
            if formula_sheet.max_row > 10_000 or formula_sheet.max_column > 256:
                raise LegacyArchiveError("历史工作表范围超过大小上限")
            cells: list[dict[str, str | None]] = []
            for formula_row, cache_row in zip(
                formula_sheet.iter_rows(), cache_sheet.iter_rows(), strict=True
            ):
                for formula_cell, cache_cell in zip(formula_row, cache_row, strict=True):
                    if formula_cell.value is None:
                        continue
                    value_type, value = _encoded_value(formula_cell.value)
                    if formula_cell.data_type == "f":
                        value_type = "formula"
                    cached_type, cached_value = _encoded_value(cache_cell.value)
                    cells.append(
                        {
                            "address": formula_cell.coordinate,
                            "type": value_type,
                            "value": value,
                            "cached_type": cached_type if value_type == "formula" else None,
                            "cached_value": cached_value if value_type == "formula" else None,
                        }
                    )
                    total_cells += 1
                    if total_cells > _MAX_CELLS:
                        raise LegacyArchiveError("历史工作簿单元格超过大小上限")
            if not cells:
                raise LegacyArchiveError(f"历史月份 {period} 没有可保存的单元格")
            cells_json = json.dumps(cells, ensure_ascii=False, separators=(",", ":"))
            sheets.append(
                ArchivedSheet(
                    period=period,
                    sheet_name=name,
                    row_count=formula_sheet.max_row,
                    cell_count=len(cells),
                    content_sha256=hashlib.sha256(cells_json.encode("utf-8")).hexdigest(),
                    cells_json=cells_json,
                )
            )
        return ArchivedWorkbook(
            source_sha256=hashlib.sha256(source_bytes).hexdigest(),
            source_bytes=source_bytes,
            original_filename=filename,
            sheets=tuple(sheets),
        )
    finally:
        formulas.close()
        cached.close()


class DatabaseLegacyArchiveStore:
    """Worker-role adapter; the caller owns one all-or-nothing transaction."""

    def __init__(self, session: Session) -> None:
        self._session = session

    def persist(
        self,
        *,
        entity_ref: UUID,
        scope_pairs: tuple[tuple[UUID, str], ...],
        actor_ref: str,
        archive: ArchivedWorkbook,
    ) -> UUID:
        if (
            entity_ref not in {item[0] for item in scope_pairs}
            or len(set(scope_pairs)) != len(scope_pairs)
            or not (1 <= len(scope_pairs) <= 32)
            or any(not ref or len(ref) > 100 for _, ref in scope_pairs)
        ):
            raise LegacyArchiveError("历史对账读取范围无效")
        for scope_entity, unit_ref in scope_pairs:
            valid = self._session.execute(
                text(
                    "SELECT 1 FROM public.business_unit "
                    "WHERE entity_id = :entity_ref AND ref = :unit_ref"
                ),
                {"entity_ref": scope_entity, "unit_ref": unit_ref},
            ).scalar_one_or_none()
            if valid is None:
                raise LegacyArchiveError("历史对账读取范围对应的门店不存在")
        scope_json = json.dumps(
            [
                {"entity_ref": str(scope_entity), "business_unit_ref": unit_ref}
                for scope_entity, unit_ref in sorted(
                    scope_pairs, key=lambda item: (str(item[0]), item[1])
                )
            ],
            separators=(",", ":"),
        )
        if hashlib.sha256(archive.source_bytes).hexdigest() != archive.source_sha256:
            raise LegacyArchiveError("历史对账来源摘要与文件内容不一致")
        if not archive.sheets or len({sheet.period for sheet in archive.sheets}) != len(
            archive.sheets
        ):
            raise LegacyArchiveError("历史对账月份清单无效")
        for sheet in archive.sheets:
            if hashlib.sha256(sheet.cells_json.encode("utf-8")).hexdigest() != sheet.content_sha256:
                raise LegacyArchiveError("历史对账单元格摘要不一致")
        digest = bytes.fromhex(archive.source_sha256)
        existing = (
            self._session.execute(
                text(
                    "SELECT source_ref, entity_ref, scope_pairs, sheet_count "
                    "FROM reconciliation_legacy.source "
                    "WHERE source_sha256 = :digest"
                ),
                {"digest": digest},
            )
            .mappings()
            .one_or_none()
        )
        if existing is not None:
            if (
                existing["entity_ref"] != entity_ref
                or existing["scope_pairs"] != json.loads(scope_json)
                or existing["sheet_count"] != len(archive.sheets)
            ):
                raise LegacyArchiveError("相同来源摘要与已有入库范围冲突")
            stored = (
                self._session.execute(
                    text(
                        "SELECT period, content_sha256, cell_count "
                        "FROM reconciliation_legacy.sheet WHERE source_ref = :source_ref"
                    ),
                    {"source_ref": existing["source_ref"]},
                )
                .mappings()
                .all()
            )
            expected = {
                (sheet.period, bytes.fromhex(sheet.content_sha256), sheet.cell_count)
                for sheet in archive.sheets
            }
            actual = {
                (row["period"], bytes(row["content_sha256"]), row["cell_count"]) for row in stored
            }
            if actual != expected:
                raise LegacyArchiveError("已有历史对账入库明细与来源摘要不一致")
            return UUID(str(existing["source_ref"]))
        source_ref = self._session.execute(
            text(
                """
                INSERT INTO reconciliation_legacy.source (
                    entity_ref, scope_pairs, source_sha256, source_bytes, original_filename,
                    sheet_count, imported_by
                ) VALUES (:entity_ref, CAST(:scope_pairs AS jsonb), :digest,
                          :source_bytes, :filename, :sheet_count, :actor_ref)
                ON CONFLICT (source_sha256) DO NOTHING
                RETURNING source_ref
                """
            ),
            {
                "entity_ref": entity_ref,
                "scope_pairs": scope_json,
                "digest": digest,
                "source_bytes": archive.source_bytes,
                "filename": archive.original_filename,
                "sheet_count": len(archive.sheets),
                "actor_ref": actor_ref,
            },
        ).scalar_one_or_none()
        if source_ref is None:
            # Another import committed the same bytes while this transaction waited
            # on the unique constraint. Re-read its immutable scope and sheet hashes.
            return self.persist(
                entity_ref=entity_ref,
                scope_pairs=scope_pairs,
                actor_ref=actor_ref,
                archive=archive,
            )
        for sheet in archive.sheets:
            self._session.execute(
                text(
                    """
                    INSERT INTO reconciliation_legacy.sheet (
                        source_ref, period, sheet_name, row_count, cell_count,
                        content_sha256, cells
                    ) VALUES (:source_ref, :period, :sheet_name, :row_count,
                              :cell_count, :content_sha256, CAST(:cells_json AS jsonb))
                    """
                ),
                {
                    "source_ref": source_ref,
                    "period": sheet.period,
                    "sheet_name": sheet.sheet_name,
                    "row_count": sheet.row_count,
                    "cell_count": sheet.cell_count,
                    "content_sha256": bytes.fromhex(sheet.content_sha256),
                    "cells_json": sheet.cells_json,
                },
            )
        return UUID(str(source_ref))
