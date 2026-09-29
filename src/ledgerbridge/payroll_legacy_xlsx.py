"""Read-only Adapter from a calculated legacy wage workbook to a payroll draft."""

# ruff: noqa: RUF001

from __future__ import annotations

import hashlib
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]

from ledgerbridge.payroll_module import (
    EmployeeType,
    PaymentChannel,
    PayrollComponent,
    PayrollDraft,
    PayrollDraftLine,
    compact_remittance_memo,
)

_FILENAME_PERIOD = re.compile(r"^(?P<year>\d{2})\.(?P<month>1[0-2]|0?[1-9])$")
_LOCATIONS = frozenset({"星汇", "薇旭", "雅阁", "逸豪", "同富", "青居客", "一品", "粥店"})
_TYPE_MAP = {
    "正式工": EmployeeType.REGULAR,
    "临时工": EmployeeType.TEMPORARY,
    "代领工资人": EmployeeType.PAYEE_ONLY,
}
_COMPONENTS = {
    "基本工资（元）": "BASE_WAGE",
    "岗位补贴（元）": "POSITION_ALLOWANCE",
    "卫生补贴（元）": "HYGIENE_ALLOWANCE",
    "夜班（元）": "NIGHT_SHIFT",
    "技术补贴（元）": "SKILL_ALLOWANCE",
    "出行补贴（元）": "TRAVEL_ALLOWANCE",
    "通讯补贴（元）": "COMMUNICATION_ALLOWANCE",
    "住宿（元）": "LODGING_ALLOWANCE",
    "餐补（元）": "MEAL_ALLOWANCE",
    "化妆补贴（元）": "MAKEUP_ALLOWANCE",
    "全勤（元）": "ATTENDANCE_ALLOWANCE",
    "请假（元）": "LEAVE_ADJUSTMENT",
    "加班（元）": "OVERTIME",
    "考核（元）": "ASSESSMENT",
    "社保（元）": "SOCIAL_SECURITY",
    "代扣（元）": "WITHHOLDING",
    "好评（元）": "PERFORMANCE",
    "特殊项金额（元）": "SPECIAL_ITEM",
}


class LegacyPayrollError(ValueError):
    """The calculated workbook is incomplete or internally inconsistent."""


def _minor(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise LegacyPayrollError(f"工资金额不是数字: {value!r}")
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _headers(sheet: Any, row: int = 2) -> dict[str, int]:
    return {
        str(sheet.cell(row, column).value).strip(): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(row, column).value is not None
    }


def _employee_id(name: str) -> str:
    return f"legacy-{hashlib.sha256(name.encode('utf-8')).hexdigest()[:20]}"


def _period(path: Path) -> str:
    match = _FILENAME_PERIOD.fullmatch(path.stem)
    if match is None:
        raise LegacyPayrollError("工资主表文件名必须为 YY.M.xlsx")
    return f"20{match['year']}-{int(match['month']):02d}"


def _require_sheets(workbook: Any) -> None:
    required = {"工资发放表", "工资明细表", "员工信息表", "考勤表"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise LegacyPayrollError(f"工资主表缺少工作表: {','.join(sorted(missing))}")


def import_calculated_payroll(
    workbook_path: str | Path,
    *,
    company_id: str,
    rules_version: str,
    cash_splits_minor: dict[str, int] | None = None,
    supplemental_minor: dict[str, int] | None = None,
) -> PayrollDraft:
    """Build a draft from cached values without modifying the source workbook."""

    path = Path(workbook_path)
    period = _period(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        _require_sheets(workbook)
        info_sheet = workbook["员工信息表"]
        info_headers = _headers(info_sheet)
        info: dict[str, dict[str, Any]] = {}
        for row in range(3, info_sheet.max_row + 1):
            name = str(info_sheet.cell(row, info_headers["姓名"]).value or "").strip()
            if name:
                info[name] = {
                    "account": info_sheet.cell(row, info_headers["卡号"]).value,
                    "alias": info_sheet.cell(row, info_headers["别名"]).value,
                    "type": info_sheet.cell(row, info_headers["人员类型"]).value,
                    "cash": info_sheet.cell(row, info_headers["是否现金发放"]).value,
                    "chen_team": info_sheet.cell(row, info_headers["是否陈队"]).value,
                }

        attendance_sheet = workbook["考勤表"]
        attendance_headers = _headers(attendance_sheet)
        attendance = {
            str(attendance_sheet.cell(row, attendance_headers["姓名"]).value).strip(): str(
                attendance_sheet.cell(row, attendance_headers["考勤天数（天）"]).value or ""
            ).strip()
            for row in range(3, attendance_sheet.max_row + 1)
            if attendance_sheet.cell(row, attendance_headers["姓名"]).value
        }

        release_sheet = workbook["工资发放表"]
        release_headers = _headers(release_sheet)
        release: dict[str, tuple[str, str, str]] = {}
        for row in range(3, release_sheet.max_row + 1):
            name = str(release_sheet.cell(row, release_headers["姓名"]).value or "").strip()
            if not name:
                continue
            release[name] = (
                str(release_sheet.cell(row, release_headers["实领人"]).value or "").strip(),
                str(release_sheet.cell(row, release_headers["卡号"]).value or "").strip(),
                str(release_sheet.cell(row, release_headers["备注"]).value or "").strip(),
            )

        detail_sheet = workbook["工资明细表"]
        detail_headers = _headers(detail_sheet)
        lines: list[PayrollDraftLine] = []
        for row in range(3, detail_sheet.max_row + 1):
            name = str(detail_sheet.cell(row, detail_headers["姓名"]).value or "").strip()
            if not name or name == "总计":
                continue
            net = _minor(detail_sheet.cell(row, detail_headers["应发"]).value)
            if net <= 0:
                continue
            location = str(
                detail_sheet.cell(row, detail_headers["工作地点"]).value or ""
            ).strip()
            if location not in _LOCATIONS:
                raise LegacyPayrollError(f"{name} 的工作地点无法生成代发表: {location}")
            profile = info.get(name)
            if profile is None:
                raise LegacyPayrollError(f"员工信息表缺少: {name}")
            employee_type = _TYPE_MAP.get(str(profile["type"] or "").strip())
            if employee_type is None:
                raise LegacyPayrollError(f"{name} 的人员类型无效")
            payee, account, memo = release.get(name, ("", "", ""))
            if not payee:
                payee = str(profile["alias"] or name).strip()
            if not account:
                account = str(profile["account"] or "").strip()
            cash = net if str(profile["cash"] or "").strip() == "是" else 0
            split = (cash_splits_minor or {}).get(name, 0)
            if cash == 0 and 0 < split < net:
                cash = split
            channel = PaymentChannel.CASH if cash == net else PaymentChannel.MYBANK
            bank = net - cash
            supplemental = (supplemental_minor or {}).get(name, 0)
            if supplemental < 0 or supplemental > bank:
                raise LegacyPayrollError(f"{name} 的补发金额超过银行代发金额")
            component_list = [
                PayrollComponent(
                    code=code,
                    amount_minor=_minor(detail_sheet.cell(row, detail_headers[header]).value),
                )
                for header, code in _COMPONENTS.items()
            ]
            if str(profile["chen_team"] or "").strip() == "是":
                component_list.append(
                    PayrollComponent(code="CHEN_TEAM_OFFSET", amount_minor=-450_000)
                )
            components = tuple(component_list)
            if sum(item.amount_minor for item in components) != net:
                raise LegacyPayrollError(f"{name} 的工资组成与应发不一致")
            lines.append(
                PayrollDraftLine(
                    employee_id=_employee_id(name),
                    employee_name=name,
                    employee_type=employee_type,
                    location=location,
                    job_group=str(
                        detail_sheet.cell(row, detail_headers["岗位"]).value or ""
                    ).strip(),
                    attendance_days=attendance.get(name, ""),
                    payment_channel=channel,
                    payee_name=payee,
                    account_number=account,
                    memo=compact_remittance_memo(memo, bank),
                    cash_amount_minor=cash,
                    supplemental_amount_minor=supplemental,
                    components=components,
                )
            )
        if not lines:
            raise LegacyPayrollError("工资主表没有可导入的已计算工资")
        return PayrollDraft(
            company_id=company_id,
            period=period,
            revision=1,
            rules_version=rules_version,
            lines=tuple(lines),
        )
    finally:
        workbook.close()
