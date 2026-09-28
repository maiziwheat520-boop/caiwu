"""Read-only Adapter for preflighting the legacy hotel reconciliation workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]

_PERIOD = re.compile(r"^(?P<year>\d{2})\.(?P<month>1[0-2]|0?[1-9])$")
_ORIGIN_PERIOD = re.compile(r"(?P<year>\d{2})\.(?P<month>1[0-2]|0?[1-9])")


class LegacyReconciliationError(ValueError):
    """The workbook cannot be migrated without an explicit correction."""


@dataclass(frozen=True, slots=True)
class LegacyAdjustment:
    origin_period: str
    effective_month: str
    source_location: str
    original_amount_minor: int
    corrected_amount_minor: int
    adjustment_amount_minor: int
    reason: str


@dataclass(frozen=True, slots=True)
class LegacyReconciliationPreflight:
    sheet_name: str
    period: str
    income_amount_minor: int
    expense_amount_minor: int
    opening_balance_minor: int
    closing_balance_minor: int
    bridge_amount_minor: int
    adjustments: tuple[LegacyAdjustment, ...]
    checks: tuple[str, ...]


def _period(sheet_name: str) -> str:
    match = _PERIOD.fullmatch(sheet_name.strip())
    if match is None:
        raise LegacyReconciliationError(f"工作表名称不是月份: {sheet_name}")
    return f"20{match['year']}-{int(match['month']):02d}"


def _minor(value: Any, *, cell: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise LegacyReconciliationError(f"{cell} 不是可用金额")
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _find_row(sheet: Any, label: str, *, column: int, start: int = 1) -> int:
    for row in range(start, sheet.max_row + 1):
        if str(sheet.cell(row, column).value or "").strip() == label:
            return row
    raise LegacyReconciliationError(f"找不到控制项: {label}")


def _read_adjustments(sheet: Any, period: str) -> tuple[LegacyAdjustment, ...]:
    header_row = _find_row(sheet, "序号", column=1)
    result: list[LegacyAdjustment] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        sequence = sheet.cell(row, 1).value
        if not isinstance(sequence, (int, float)):
            continue
        source = str(sheet.cell(row, 2).value or "").strip()
        match = _ORIGIN_PERIOD.search(source)
        if match is None:
            raise LegacyReconciliationError(f"B{row} 缺少原始月份")
        result.append(
            LegacyAdjustment(
                origin_period=f"20{match['year']}-{int(match['month']):02d}",
                effective_month=period,
                source_location=source,
                original_amount_minor=_minor(sheet.cell(row, 3).value, cell=f"C{row}"),
                corrected_amount_minor=_minor(sheet.cell(row, 4).value, cell=f"D{row}"),
                adjustment_amount_minor=_minor(sheet.cell(row, 5).value, cell=f"E{row}"),
                reason=str(sheet.cell(row, 6).value or "").strip(),
            )
        )
    if not result:
        raise LegacyReconciliationError("历史差错调整明细为空")
    return tuple(result)


def preflight_legacy_reconciliation(
    workbook_path: str | Path, sheet_name: str
) -> LegacyReconciliationPreflight:
    """Read cached workbook values and prove the sheet's closing bridge."""

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise LegacyReconciliationError(f"工作表不存在: {sheet_name}")
        sheet = workbook[sheet_name]
        period = _period(sheet_name)
        expense_heading = _find_row(sheet, "支出明细", column=1)
        income_row = _find_row(sheet, "收入共计", column=1)
        expense_row = _find_row(sheet, "支出共计", column=1, start=expense_heading)
        opening_row = _find_row(sheet, "结余", column=1, start=expense_row)
        closing_row = _find_row(sheet, "共计", column=1, start=opening_row + 1)

        income = _minor(sheet.cell(income_row, 2).value, cell=f"B{income_row}")
        expense = _minor(sheet.cell(expense_row, 2).value, cell=f"B{expense_row}")
        opening = _minor(sheet.cell(opening_row, 2).value, cell=f"B{opening_row}")
        closing = _minor(sheet.cell(closing_row, 2).value, cell=f"B{closing_row}")
        bridge = sum(
            _minor(value, cell=f"{column}{row}")
            for row in range(opening_row, closing_row)
            for column in ("B", "E")
            if (value := sheet[f"{column}{row}"].value) is not None
        )
        if bridge != closing:
            raise LegacyReconciliationError(
                f"期末桥接不平: 计算 {bridge} 分, 表内 {closing} 分"
            )

        adjustments = _read_adjustments(sheet, period)
        history_row = _find_row(sheet, "历史差错调整", column=4, start=opening_row)
        history_amount = abs(_minor(sheet.cell(history_row, 5).value, cell=f"E{history_row}"))
        if sum(item.adjustment_amount_minor for item in adjustments) != history_amount:
            raise LegacyReconciliationError("历史差错调整明细与结余区不一致")
        return LegacyReconciliationPreflight(
            sheet_name=sheet_name,
            period=period,
            income_amount_minor=income,
            expense_amount_minor=expense,
            opening_balance_minor=opening,
            closing_balance_minor=closing,
            bridge_amount_minor=bridge,
            adjustments=adjustments,
            checks=("CONTROL_TOTALS_FOUND", "CLOSING_BRIDGE_BALANCED", "ADJUSTMENTS_BALANCED"),
        )
    finally:
        workbook.close()
